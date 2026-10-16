"""
Relatório de Criptomoedas em Tempo Real
-----------------------------------------
Busca a cotação atual de criptomoedas na API pública da CoinGecko
(preço em reais e em dólares, variação de 24h e 7 dias, volume,
capitalização de mercado e distância da máxima histórica), calcula
um resumo e gera três saídas: Excel formatado com gráfico, PDF e
resumo no terminal.

Não precisa de chave de API nem cadastro: é um endpoint público e gratuito.

Uso:
    python relatorio_cripto.py                       -> usa 10 moedas padrão
    python relatorio_cripto.py bitcoin ethereum solana -> usa só as moedas informadas

Autor: Paulo Roberto
"""

import argparse
import os
import sys
from datetime import datetime

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from fpdf import FPDF


# ----------------------------------------------------------------------
# 1. CONFIGURAÇÃO
# ----------------------------------------------------------------------

MOEDAS_PADRAO = [
    "bitcoin", "ethereum", "binancecoin", "solana", "cardano",
    "ripple", "dogecoin", "polkadot", "litecoin", "chainlink",
]

URL_API = "https://api.coingecko.com/api/v3/coins/markets"


# ----------------------------------------------------------------------
# 2. ARGUMENTOS DE LINHA DE COMANDO
# ----------------------------------------------------------------------

def ler_argumentos():
    parser = argparse.ArgumentParser(description="Gera um relatório de criptomoedas em tempo real.")
    parser.add_argument(
        "moedas",
        nargs="*",
        help="IDs das moedas no CoinGecko (ex: bitcoin ethereum). Se não informar nada, usa uma lista padrão de 10 moedas conhecidas.",
    )
    return parser.parse_args()


# ----------------------------------------------------------------------
# 3. BUSCA DOS DADOS NA API
# ----------------------------------------------------------------------

def buscar_precos(moedas, moeda_destino):
    """Busca os dados de mercado para uma lista de moedas, numa moeda de destino (brl ou usd)."""
    parametros = {
        "vs_currency": moeda_destino,
        "ids": ",".join(moedas),
        "order": "market_cap_desc",
        "price_change_percentage": "24h,7d",
    }
    resposta = requests.get(URL_API, params=parametros, timeout=10)
    resposta.raise_for_status()
    return resposta.json()


def buscar_cotacoes(moedas):
    """
    Busca os dados em reais (completo) e em dólares (só o preço, pra conversão),
    e devolve None se der algum problema de conexão.
    """
    try:
        dados_brl = buscar_precos(moedas, "brl")
        dados_usd = buscar_precos(moedas, "usd")
    except requests.exceptions.ConnectionError:
        print("Não foi possível conectar à internet. Verifique sua conexão e tente novamente.")
        return None
    except requests.exceptions.Timeout:
        print("A API demorou demais para responder. Tente novamente em instantes.")
        return None
    except requests.exceptions.HTTPError as erro:
        print(f"A API respondeu com um erro: {erro}")
        return None
    except requests.exceptions.RequestException as erro:
        print(f"Erro inesperado ao buscar os dados: {erro}")
        return None

    if not dados_brl:
        print("Nenhuma moeda encontrada com os IDs informados. Confira se os nomes estão corretos (ex: bitcoin, ethereum).")
        return None

    mapa_usd = {m["id"]: m["current_price"] for m in dados_usd}
    for moeda in dados_brl:
        moeda["preco_usd"] = mapa_usd.get(moeda["id"])

    return dados_brl


# ----------------------------------------------------------------------
# 4. PROCESSAMENTO DOS DADOS
# ----------------------------------------------------------------------

def classificar_tendencia(var_24h, var_7d):
    """
    Classificação simples baseada no que já aconteceu (não é previsão):
    compara a direção da variação de 24h com a de 7 dias.
    """
    if var_24h > 0 and var_7d > 0:
        return "Em alta"
    if var_24h < 0 and var_7d < 0:
        return "Em queda"
    return "Sinal misto"


def processar_dados(moedas_brutas):
    """Transforma a resposta bruta da API numa lista de dicionários mais simples de usar."""
    processadas = []
    for m in moedas_brutas:
        var_24h = m.get("price_change_percentage_24h") or 0
        var_7d = m.get("price_change_percentage_7d_in_currency") or 0

        processadas.append({
            "nome": m["name"],
            "simbolo": m["symbol"].upper(),
            "preco_brl": m["current_price"],
            "preco_usd": m["preco_usd"],
            "var_24h": var_24h,
            "var_7d": var_7d,
            "tendencia": classificar_tendencia(var_24h, var_7d),
            "maxima_24h": m.get("high_24h") or 0,
            "minima_24h": m.get("low_24h") or 0,
            "market_cap": m.get("market_cap") or 0,
            "volume_24h": m.get("total_volume") or 0,
            "dist_ath": m.get("ath_change_percentage") or 0,
        })
    return processadas


# ----------------------------------------------------------------------
# 5. CÁLCULO DO RESUMO
# ----------------------------------------------------------------------

def montar_resumo(moedas):
    por_variacao = sorted(moedas, key=lambda m: m["var_24h"], reverse=True)
    maior_alta = por_variacao[0]
    maior_queda = por_variacao[-1]
    mais_cara = max(moedas, key=lambda m: m["preco_brl"])
    maior_cap = max(moedas, key=lambda m: m["market_cap"])
    maior_volume = max(moedas, key=lambda m: m["volume_24h"])
    mais_perto_do_ath = max(moedas, key=lambda m: m["dist_ath"])  # dist_ath é negativo; o menos negativo é o mais perto
    variacao_media = sum(m["var_24h"] for m in moedas) / len(moedas)

    metricas = [
        ("Maior alta em 24h", f"{maior_alta['nome']} ({maior_alta['var_24h']:+.2f}%)"),
        ("Maior queda em 24h", f"{maior_queda['nome']} ({maior_queda['var_24h']:+.2f}%)"),
        ("Moeda de maior valor unitário", f"{mais_cara['nome']} (R$ {mais_cara['preco_brl']:,.2f})"),
        ("Maior capitalização de mercado", f"{maior_cap['nome']}"),
        ("Maior volume negociado em 24h", f"{maior_volume['nome']}"),
        ("Mais próxima da máxima histórica", f"{mais_perto_do_ath['nome']} ({mais_perto_do_ath['dist_ath']:+.1f}% do ATH)"),
        ("Variação média do grupo (24h)", f"{variacao_media:+.2f}%"),
    ]

    ranking = [(m["nome"], m["var_24h"], m["var_7d"], m["tendencia"]) for m in por_variacao]

    return {
        "titulo": "Relatório de Criptomoedas",
        "metricas": metricas,
        "ranking_titulo": "Variação por moeda (24h e 7 dias)",
        "ranking": ranking,
    }


# ----------------------------------------------------------------------
# 6. GERAÇÃO DO EXCEL FORMATADO
# ----------------------------------------------------------------------

COLUNAS_DADOS = [
    "Nome", "Símbolo", "Preço (R$)", "Preço (US$)", "Variação 24h (%)", "Variação 7d (%)",
    "Tendência", "Máxima 24h (R$)", "Mínima 24h (R$)", "Market Cap (R$)", "Volume 24h (R$)",
    "Distância do ATH (%)",
]


def gerar_excel(moedas, resumo, caminho_saida):
    wb = Workbook()

    aba_dados = wb.active
    aba_dados.title = "Dados"
    aba_dados.append(COLUNAS_DADOS)

    for m in moedas:
        aba_dados.append([
            m["nome"], m["simbolo"],
            round(m["preco_brl"], 2), round(m["preco_usd"], 2),
            round(m["var_24h"], 2), round(m["var_7d"], 2),
            m["tendencia"],
            round(m["maxima_24h"], 2), round(m["minima_24h"], 2),
            round(m["market_cap"], 2), round(m["volume_24h"], 2),
            round(m["dist_ath"], 2),
        ])

    estilo_cabecalho(aba_dados, len(COLUNAS_DADOS))

    # Colore a coluna de variação 24h (5) e 7d (6) conforme sinal
    for linha in range(2, aba_dados.max_row + 1):
        for col in (5, 6):
            celula = aba_dados.cell(row=linha, column=col)
            if isinstance(celula.value, (int, float)):
                celula.font = Font(color="1F7A3D", bold=True) if celula.value >= 0 else Font(color="C0392B", bold=True)

    ajustar_largura_colunas(aba_dados)
    adicionar_grafico(aba_dados, len(moedas))

    # --- Aba 2: resumo ---
    aba_resumo = wb.create_sheet("Resumo")
    aba_resumo["A1"] = resumo["titulo"]
    aba_resumo["A1"].font = Font(size=16, bold=True, color="1F4E78")
    timestamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    aba_resumo["A2"] = f"Gerado em {timestamp}"
    aba_resumo["A2"].font = Font(italic=True, size=9, color="808080")

    linha_atual = 4
    aba_resumo.cell(row=linha_atual, column=1, value="Indicador").font = Font(bold=True, color="FFFFFF")
    aba_resumo.cell(row=linha_atual, column=2, value="Valor").font = Font(bold=True, color="FFFFFF")
    for col in (1, 2):
        aba_resumo.cell(row=linha_atual, column=col).fill = PatternFill("solid", fgColor="1F4E78")
    linha_atual += 1

    for nome, valor in resumo["metricas"]:
        aba_resumo.cell(row=linha_atual, column=1, value=nome)
        aba_resumo.cell(row=linha_atual, column=2, value=valor)
        linha_atual += 1

    linha_atual += 1
    aba_resumo.cell(row=linha_atual, column=1, value=resumo["ranking_titulo"]).font = Font(bold=True, size=12)
    linha_atual += 1

    cabecalho_ranking = ["Moeda", "Var. 24h (%)", "Var. 7d (%)", "Tendência"]
    for col, texto in enumerate(cabecalho_ranking, start=1):
        aba_resumo.cell(row=linha_atual, column=col, value=texto).font = Font(bold=True)
    linha_atual += 1

    for nome, var24, var7, tendencia in resumo["ranking"]:
        cor = "1F7A3D" if var24 >= 0 else "C0392B"
        aba_resumo.cell(row=linha_atual, column=1, value=nome)
        aba_resumo.cell(row=linha_atual, column=2, value=round(var24, 2)).font = Font(color=cor, bold=True)
        aba_resumo.cell(row=linha_atual, column=3, value=round(var7, 2)).font = Font(color=cor)
        aba_resumo.cell(row=linha_atual, column=4, value=tendencia)
        linha_atual += 1

    linha_atual += 2
    aviso = ("Fonte: CoinGecko API (api.coingecko.com). Consulta realizada em " + timestamp +
             ". Relatório apenas informativo, não constitui recomendação de investimento.")
    aba_resumo.cell(row=linha_atual, column=1, value=aviso).font = Font(italic=True, size=8, color="808080")

    ajustar_largura_colunas(aba_resumo)
    wb.save(caminho_saida)


def adicionar_grafico(aba_dados, n_moedas):
    """Cria um gráfico de barras nativo do Excel comparando a variação de 24h entre as moedas."""
    grafico = BarChart()
    grafico.title = "Variação em 24h por moeda"
    grafico.y_axis.title = "% de variação"
    grafico.x_axis.title = "Moeda"
    grafico.height = 8
    grafico.width = 18

    dados_ref = Reference(aba_dados, min_col=5, min_row=1, max_row=1 + n_moedas, max_col=5)
    categorias_ref = Reference(aba_dados, min_col=1, min_row=2, max_row=1 + n_moedas)
    grafico.add_data(dados_ref, titles_from_data=True)
    grafico.set_categories(categorias_ref)

    aba_dados.add_chart(grafico, f"B{n_moedas + 4}")


def estilo_cabecalho(aba, n_colunas):
    for col in range(1, n_colunas + 1):
        cel = aba.cell(row=1, column=col)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(horizontal="center")


def ajustar_largura_colunas(aba):
    for col_cells in aba.columns:
        comprimento = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        letra = get_column_letter(col_cells[0].column)
        aba.column_dimensions[letra].width = min(comprimento + 4, 30)


# ----------------------------------------------------------------------
# 7. GERAÇÃO DO PDF
# ----------------------------------------------------------------------

def gerar_pdf(resumo, timestamp, caminho_saida):
    pdf = FPDF()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(31, 78, 120)
    pdf.cell(0, 12, resumo["titulo"], ln=True)

    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(128, 128, 128)
    pdf.cell(0, 8, f"Gerado em {timestamp}", ln=True)
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(31, 78, 120)
    pdf.cell(0, 10, "Principais indicadores", ln=True)
    pdf.set_draw_color(31, 78, 120)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    for nome, valor in resumo["metricas"]:
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(30, 30, 30)
        pdf.cell(85, 7, nome)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 7, str(valor), ln=True)

    pdf.ln(6)
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(31, 78, 120)
    pdf.cell(0, 10, resumo["ranking_titulo"], ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(70, 6, "Moeda")
    pdf.cell(35, 6, "Var. 24h")
    pdf.cell(35, 6, "Var. 7d")
    pdf.cell(0, 6, "Tendência", ln=True)

    pdf.set_font("Helvetica", "", 9)
    for nome, var24, var7, tendencia in resumo["ranking"]:
        cor = (31, 122, 61) if var24 >= 0 else (192, 57, 43)
        pdf.set_text_color(*cor)
        pdf.cell(70, 6, str(nome))
        pdf.cell(35, 6, f"{var24:+.2f}%")
        pdf.cell(35, 6, f"{var7:+.2f}%")
        pdf.set_text_color(60, 60, 60)
        pdf.cell(0, 6, tendencia, ln=True)

    pdf.ln(8)
    pdf.set_font("Helvetica", "I", 7)
    pdf.set_text_color(130, 130, 130)
    pdf.multi_cell(0, 4, (
        "Fonte: CoinGecko API (api.coingecko.com). Consulta realizada em " + timestamp +
        ". Relatório apenas informativo, não constitui recomendação de investimento."
    ))

    pdf.output(caminho_saida)


# ----------------------------------------------------------------------
# 8. RESUMO NO TERMINAL
# ----------------------------------------------------------------------

def mostrar_resumo_terminal(resumo):
    print("\n" + "=" * 60)
    print(resumo["titulo"])
    print("=" * 60)
    for nome, valor in resumo["metricas"]:
        print(f"- {nome}: {valor}")
    print("-" * 60)
    print(resumo["ranking_titulo"] + ":")
    for nome, var24, var7, tendencia in resumo["ranking"]:
        print(f"  {nome}: 24h {var24:+.2f}% | 7d {var7:+.2f}% | {tendencia}")
    print("=" * 60 + "\n")


# ----------------------------------------------------------------------
# 9. PROGRAMA PRINCIPAL
# ----------------------------------------------------------------------

def main():
    args = ler_argumentos()
    moedas_escolhidas = args.moedas if args.moedas else MOEDAS_PADRAO

    print(f"Buscando cotações de {len(moedas_escolhidas)} moeda(s)... aguarde um instante.")
    dados_brutos = buscar_cotacoes(moedas_escolhidas)

    if dados_brutos is None:
        sys.exit(1)

    moedas = processar_dados(dados_brutos)
    resumo = montar_resumo(moedas)
    mostrar_resumo_terminal(resumo)

    pasta_atual = os.path.dirname(os.path.abspath(__file__))
    pasta_saida = os.path.join(pasta_atual, "saida")
    os.makedirs(pasta_saida, exist_ok=True)

    agora = datetime.now()
    timestamp = agora.strftime("%d/%m/%Y %H:%M")
    sufixo_arquivo = agora.strftime("%Y%m%d_%H%M")

    caminho_excel = os.path.join(pasta_saida, f"relatorio_cripto_{sufixo_arquivo}.xlsx")
    caminho_pdf = os.path.join(pasta_saida, f"relatorio_cripto_{sufixo_arquivo}.pdf")

    gerar_excel(moedas, resumo, caminho_excel)
    gerar_pdf(resumo, timestamp, caminho_pdf)

    print(f"Excel gerado em: {caminho_excel}")
    print(f"PDF gerado em: {caminho_pdf}")


if __name__ == "__main__":
    main()
